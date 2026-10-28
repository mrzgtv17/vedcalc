"""Обновление вшитого справочника ТН ВЭД ЕАЭС (коды, наименования, ставки ввозной пошлины).

Источник — бесплатная ежедневная выгрузка tws.by («как есть, без гарантий»).
Скачивает Excel, проверяет структуру и пишет vedcalc/data/tnved_tws.json — по строке на код,
чтобы изменения хорошо читались в git. Печатает сводку отличий от прежней версии.

Использование: python scripts/update_tnved.py [--xlsx путь_к_файлу]
"""

from __future__ import annotations

import argparse
import json
import sys
import tempfile
import urllib.request
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
TARGET = ROOT / "vedcalc" / "data" / "tnved_tws.json"
SOURCE_URL = "https://www.tws.by/tws/tnved/download/excel"
DETAILS_URL = "https://www.tws.by/tws/tnved/code/{id}"  # id — внутренний номер tws.by
MIN_CODES = 10_000  # меньше — значит, выгрузка сломалась
ARROW = " \U0001f83a "  # разделитель уровней наименования в выгрузке tws.by


def download(path: Path) -> None:
    sys.path.insert(0, str(ROOT))
    from vedcalc.core.fx import _ssl_context

    request = urllib.request.Request(SOURCE_URL, headers={"User-Agent": "vedcalc-update"})
    with urllib.request.urlopen(request, timeout=300, context=_ssl_context()) as response:
        path.write_bytes(response.read())


def parse(path: Path) -> dict:
    # Не read_only: нужны гиперссылки «подробнее» (внутренние номера кодов на tws.by)
    wb = load_workbook(path, data_only=True)
    info = {row[0]: row[1] for row in wb.worksheets[0].iter_rows(values_only=True) if row[0]}
    actual = datetime.strptime(str(info["Актуальность данных"]), "%d.%m.%Y").date()
    rows = wb.worksheets[1].iter_rows()
    header = [str(c.value) for c in next(rows)][:3]
    if header != ["Код", "Наименование", "Тариф"]:
        raise SystemExit(f"Неожиданные колонки выгрузки: {header}")
    codes = {}
    for cells in rows:
        code, name, tariff = (str(c.value or "").strip() for c in cells[:3])
        if len(code) != 10 or not code.isdigit():
            continue
        tariff = "" if tariff == "None" else tariff
        link = cells[3].hyperlink.target if len(cells) > 3 and cells[3].hyperlink else ""
        tws_id = link.rstrip("/").rsplit("/", 1)[-1] if link else ""
        codes[code] = [name.replace(ARROW, " → "), tariff, tws_id]
    if len(codes) < MIN_CODES:
        raise SystemExit(f"Слишком мало кодов в выгрузке: {len(codes)}")
    return {"actual_date": actual.isoformat(), "codes": codes}


def write(data: dict) -> None:
    lines = [
        "{",
        '  "source": "tws.by — ежедневная выгрузка ТН ВЭД ЕАЭС, «как есть, без гарантий»",',
        f'  "source_url": {json.dumps(SOURCE_URL)},',
        f'  "details_url": {json.dumps(DETAILS_URL)},',
        f'  "actual_date": "{data["actual_date"]}",',
        '  "codes": {',
    ]
    items = sorted(data["codes"].items())
    for n, (code, value) in enumerate(items):
        comma = "," if n < len(items) - 1 else ""
        lines.append(f'    "{code}": {json.dumps(value, ensure_ascii=False)}{comma}')
    lines += ["  }", "}", ""]
    TARGET.write_text("\n".join(lines), encoding="utf-8")


def diff(old: dict, new: dict) -> str:
    o, n = old.get("codes", {}), new["codes"]
    added = sorted(set(n) - set(o))
    removed = sorted(set(o) - set(n))
    changed = sorted(c for c in set(o) & set(n) if o[c][1] != n[c][1])
    out = [
        f"Актуальность: {old.get('actual_date', '—')} → {new['actual_date']}",
        f"Кодов: {len(o)} → {len(n)} (новых {len(added)}, исчезло {len(removed)}, "
        f"изменилась ставка {len(changed)})",
    ]
    out += [f"  ставка {c}: {o[c][1]} → {n[c][1]}" for c in changed[:30]]
    return "\n".join(out)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--xlsx", type=Path, help="уже скачанный файл выгрузки")
    args = parser.parse_args()
    with tempfile.TemporaryDirectory() as tmp:
        path = args.xlsx or Path(tmp) / "tnved.xlsx"
        if not args.xlsx:
            download(path)
        new = parse(path)
    old = json.loads(TARGET.read_text("utf-8")) if TARGET.exists() else {}
    write(new)
    print(diff(old, new))


if __name__ == "__main__":
    main()
