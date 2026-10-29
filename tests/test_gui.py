"""GUI: контрольный пример на экране, ввод, режимы цены, файлы, экспорт, курсы."""

from datetime import date
from decimal import Decimal as D

import pytest

pytest.importorskip("PySide6")

from conftest import TEST_RATES, StaticRates
from PySide6.QtWidgets import QLabel, QMessageBox

from vedcalc.core.fx import FxError
from vedcalc.core.models import (
    DeliveryMode,
    FinancingMode,
    PriceMode,
    PricingSource,
    TaxRegime,
)
from vedcalc.core.units import Unit
from vedcalc.gui.main_window import MainWindow

NBSP = " "


def money(text: str) -> str:
    return text.replace(" ", NBSP)


@pytest.fixture
def window(qtbot, monkeypatch):
    # Диалоги «Сохранить изменения?» в тестах не показываем
    monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: QMessageBox.StandardButton.Discard)
    monkeypatch.setattr(QMessageBox, "warning", lambda *a, **k: None)
    w = MainWindow(TEST_RATES, threaded=False, extra_modules=[])
    qtbot.addWidget(w)
    return w


@pytest.fixture
def example(window):
    window.open_example()
    return window


def results(w):
    return w.calculator.results


def row(w, label):
    return w.calculator.form.items_table.rows[0][label]


def test_new_calculation_shows_hints_not_errors(window):
    status = results(window).status.text()
    assert "Заполните данные" in status
    assert window.calculator.report is None
    assert not window.pdf_action.isEnabled()


def pricing(w):
    w.tabs.setCurrentWidget(w.pricing)
    w.pricing.recalc()
    return w.pricing.results


def test_control_example_on_screen(example):
    r = results(example)
    (card,) = r.unit_cards
    # Вкладка «Импорт» — себестоимость единицы, безубыточной цены больше нет
    assert card.view.per == "за 1 шт"
    assert card.metric_text("cost") == money("147 731,00")
    assert card.metric_text("vat") == money("30 030,00")
    assert "break_even" not in card.tiles
    assert card.line_text("Таможенная стоимость") == money("130 000,00")
    assert card.line_text("Таможенный сбор") == money("1 231,00")
    assert card.line_text("Ввозная пошлина") == money("6 500,00")
    assert card.line_text("НДС при ввозе") == money("30 030,00")
    assert r.batch_view.value_text("Себестоимость партии без НДС") == money("147 731,00")
    assert r.batch_view.value_text("НДС-разрыв при P = C") is None
    text = r.details.toPlainText()
    assert "Себестоимость за 1 шт" in text and "Безубыточная" not in text
    assert example.pdf_action.isEnabled()


def test_pricing_tab_uses_import_cost(example):
    r = pricing(example)
    (card,) = r.unit_cards
    # В файле примера сохранён сценарий «продажа партии за 200 000»
    assert card.metric_text("cost") == money("147 731,00")
    assert card.metric_text("price") == money("200 000,00")
    assert card.metric_text("profit") == money("39 201,75")
    assert "35,38%" in r.target_summary.text() and "26,13%" in r.target_summary.text()
    assert r.batch_view.value_text("Цена партии с НДС") == money("244 000,00")
    assert not r.financing_view.isVisible()
    assert example.pdf_action.isEnabled()


def test_pricing_follows_import_changes(example, qtbot):
    pricing(example)
    row(example, "quantity").setText("4")
    row(example, "unit").set_value(Unit.KG)
    example.calculator.recalc()
    qtbot.waitUntil(lambda: example.pricing.report.basis.items[0].quantity == 4, timeout=2000)
    (card,) = example.pricing.results.unit_cards
    assert card.view.per == "за 1 кг"
    assert card.metric_text("cost") == money("115 990,50")


def test_unit_cards_follow_quantity_and_unit(example):
    row(example, "quantity").setText("4")
    row(example, "unit").set_value(Unit.KG)
    example.calculator.recalc()
    (card,) = results(example).unit_cards
    assert card.view.per == "за 1 кг"
    assert card.line_text("Таможенный сбор") == money("615,50")  # 2 462 / 4
    assert card.metric_text("cost") == money("115 990,50")


def test_unit_price_mode_in_gui(example):
    r = pricing(example)
    r.mode_combo.set_value(PriceMode.UNIT_PRICE)
    r.value_edit.setText("160000")
    example.pricing.recalc()
    (card,) = r.unit_cards
    assert card.metric_text("price") == money("160 000,00")
    assert "за 1 шт" in r.target_summary.text()


def test_cost_of_money_in_gui(example):
    tab = example.pricing
    r = pricing(example)
    r.mode_combo.set_value(PriceMode.MARKUP)
    r.value_edit.setText("20")
    tab.fin_mode.set_value(FinancingMode.SIMPLE)
    tab.fin_rate.setText("20")
    tab.fin_days.setText("60")
    tab.recalc()
    assert tab.report.target.financing == D("5844.20")
    assert r.unit_cards[0].metric_text("price") == money("184 290,24")
    assert r.financing_view.value_text("Стоимость денег за партию") == money("5 844,20")
    tab.fin_mode.set_value(FinancingMode.SCHEDULE)
    tab.recalc()
    assert tab.report.target.financing_detail.periods  # график посчитан
    assert "Деньги в обороте" in r.details.toPlainText() or r.financing_view.isVisible()
    # Параметры стоимости денег сохраняются в файле расчёта
    saved = example.calculator.shipment()
    assert saved.financing.mode is FinancingMode.SCHEDULE


def test_manual_cost_in_pricing_tab(example):
    tab = example.pricing
    tab.source_combo.set_value(PricingSource.MANUAL)
    tab.m_name.setText("Стеарин")
    tab.m_qty.setText("1000")
    tab.m_unit.set_value(Unit.KG)
    tab.m_cost.setText("150,5")
    tab.m_vat.setText("33")
    r = pricing(example)
    r.mode_combo.set_value(PriceMode.UNIT_PRICE)
    r.value_edit.setText("200")
    tab.recalc()
    (card,) = r.unit_cards
    assert card.view.per == "за 1 кг"
    assert card.metric_text("cost") == money("150,50")
    assert tab.report.target.gross_profit == D("49500.00")


def test_container_load_hint(example):
    form = example.calculator.form
    form.delivery_combo.set_value(DeliveryMode.DC20)
    row(example, "gross").setText("30000")
    row(example, "volume").setText("20")
    example.calculator.recalc()
    card = results(example).load_card
    assert card is not None
    assert card.bar_text("volume").startswith("60%")
    assert card.bar_text("weight").startswith("111%")
    warning = card.findChild(QLabel, "loadWarning")
    assert warning and "не помещается" in warning.text() and "2" in warning.text()
    form.delivery_combo.set_value(DeliveryMode.AIR)
    example.calculator.recalc()
    assert results(example).load_card is None


def test_editing_recalculates_on_the_fly(example, qtbot):
    row(example, "quantity").setText("2")
    row(example, "price").setText("50 000")
    qtbot.waitUntil(
        lambda: (
            example.calculator.report is not None
            and example.calculator.report.cost.items[0].quantity == 2
        ),
        timeout=2000,
    )
    report = example.calculator.report
    assert report.cost.cost_net == D("147731.00")
    assert report.cost.items[0].unit_cost_net == D("73865.50")
    assert example.dirty


def test_invalid_input_is_highlighted(example):
    example.calculator.touched = True
    row(example, "quantity").setText("abc")
    example.calculator.recalc()
    assert example.calculator.report is None
    assert "Проверьте ввод" in results(example).status.text()
    assert "не число" in row(example, "quantity").toolTip()
    assert row(example, "quantity").property("tone") == "fieldError"
    assert not example.pdf_action.isEnabled()
    row(example, "quantity").setText("1")
    example.calculator.recalc()
    assert not row(example, "quantity").property("tone")


def test_price_modes(example):
    r = pricing(example)
    r.mode_combo.set_value(PriceMode.MARGIN)
    r.value_edit.setText("25")
    example.pricing.recalc()
    assert r.unit_cards[0].metric_text("price") == money("196 974,67")
    assert r.target_summary.text().startswith("Цена с маржинальностью 25%")

    r.value_edit.setText("100")
    example.pricing.recalc()
    assert "меньше 100%" in r.target_error.text()
    assert example.calculator.report is not None  # себестоимость считается


def test_usn_regime(example):
    form = example.calculator.form
    form.regime_combo.set_value(TaxRegime.USN_INCOME_MINUS_EXPENSES)
    example.calculator.recalc()
    tab = example.pricing
    assert not tab.sale_vat.isEnabled()
    assert tab.profit_tax.text() == "15"
    assert example.calculator.report.cost.cost_net == D("177761.00")
    pricing(example)
    assert tab.report.target.sale_vat == 0


def test_manual_rate(example):
    row(example, "currency").set_value("CNY")
    row(example, "price").setText("10000")
    example.calculator.recalc()
    assert example.calculator.report.cost.items[0].goods == D("115000.00")  # 11,5 из теста
    form = example.calculator.form
    form.manual_check.setChecked(True)
    form.manual_edits["CNY"].setText("12")
    example.calculator.recalc()
    report = example.calculator.report
    assert report.cost.items[0].goods == D("120000.00")
    assert "Курс изменён вручную" in results(example).status.text()


def test_save_open_duplicate(example, tmp_path):
    path = tmp_path / "calc.vedcalc.json"
    assert example._save_to(path)
    assert not example.dirty and example.windowTitle().startswith("calc.vedcalc.json")
    example.new_file()
    assert example.calculator.report is None
    example.open_file(path)
    assert example.calculator.report.cost.cost_net == D("147731.00")
    assert example.calculator.form.name_edit.text() == "Контрольный пример"
    example.duplicate()
    assert example.calculator.form.name_edit.text() == "Контрольный пример (копия)"
    assert example.path is None and example.dirty


def test_save_refused_with_parse_errors(example, tmp_path):
    row(example, "price").setText("сто")
    assert not example._save_to(tmp_path / "x.json")
    assert not (tmp_path / "x.json").exists()


def test_exports(example, tmp_path):
    pdf, xlsx = tmp_path / "r.pdf", tmp_path / "r.xlsx"
    example.export_pdf(pdf)
    example.export_xlsx(xlsx)
    assert pdf.read_bytes()[:5] == b"%PDF-" and pdf.stat().st_size > 5000
    from openpyxl import load_workbook

    ws = load_workbook(xlsx)["Итоги"]
    values = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(1, ws.max_row + 1)}
    assert values["Себестоимость партии без НДС"] == 147731
    assert "Чистая прибыль" not in values  # на вкладке «Импорт» — только себестоимость
    # С вкладки «Цена продажи» — цена и прибыль
    pricing(example)
    example.export_pdf(pdf)
    example.export_xlsx(xlsx)
    ws = load_workbook(xlsx)["Итоги"]
    values = {ws.cell(r, 1).value: ws.cell(r, 2).value for r in range(1, ws.max_row + 1)}
    assert values["Чистая прибыль"] == 39201.75


def test_rates_load_in_background(qtbot, monkeypatch):
    monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: QMessageBox.StandardButton.Discard)
    w = MainWindow(StaticRates({"CNY": ("11.5", 1)}), threaded=True, extra_modules=[])
    qtbot.addWidget(w)
    form = w.calculator.form
    form.items_table.rows[0]["name"].setText("Насос")
    form.items_table.rows[0]["price"].setText("1000")  # CNY по умолчанию
    qtbot.waitUntil(lambda: w.calculator.report is not None, timeout=3000)
    assert w.calculator.report.cost.items[0].goods == D("11500.00")
    assert "11,5000" in form.rates_info.text()


def test_offline_without_cache_suggests_manual_rate(qtbot, monkeypatch):
    class Offline:
        def get_rates(self, on: date):
            raise FxError("Нет связи с ЦБ РФ и нет сохранённых курсов")

    monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: QMessageBox.StandardButton.Discard)
    w = MainWindow(Offline(), threaded=False, extra_modules=[])
    qtbot.addWidget(w)
    form = w.calculator.form
    form.items_table.rows[0]["name"].setText("Насос")
    form.items_table.rows[0]["price"].setText("1000")
    w.calculator.recalc()
    assert "задайте курс вручную" in results(w).status.text()
    form.manual_check.setChecked(True)
    form.manual_edits["CNY"].setText("12,5")
    w.calculator.recalc()
    assert w.calculator.report.cost.items[0].goods == D("12500.00")


def test_update_check_offers_new_version(qtbot, monkeypatch, tmp_path):
    import json

    from PySide6.QtCore import QSettings

    from vedcalc import updater
    from vedcalc.gui import updates

    data = {
        "tag_name": "v9.9.9",
        "body": "Новое",
        "html_url": "https://example.test",
        "assets": [
            {
                "name": f"vedcalc-9.9.9-macos-{updater.machine_arch()}.dmg",
                "size": 1,
                "browser_download_url": "https://example.test/a.dmg",
                "digest": "sha256:" + "0" * 64,
            }
        ],
    }
    w = MainWindow(TEST_RATES, threaded=False, extra_modules=[])
    qtbot.addWidget(w)
    w.updates.fetch = lambda url: json.dumps(data).encode()
    w.updates.settings = QSettings(str(tmp_path / "s.ini"), QSettings.Format.IniFormat)
    shown = {}

    def fake_exec(dialog):
        shown["install_enabled"] = dialog.install.isEnabled()
        dialog.choice = updates.UpdateDialog.SKIP
        return 0

    monkeypatch.setattr(updates.UpdateDialog, "exec", fake_exec)
    w.updates.check(manual=True)
    # Из исходников установка недоступна — только «скачать вручную»
    assert shown == {"install_enabled": False}
    assert w.updates.settings.value(updates.SETTINGS_SKIPPED) == "9.9.9"
    # Пропущенная версия при автоматической проверке больше не предлагается
    shown.clear()
    w.updates.check(manual=False)
    assert shown == {}


def test_update_check_reports_latest(qtbot, monkeypatch, tmp_path):
    import json

    from PySide6.QtCore import QSettings

    from vedcalc import __version__

    w = MainWindow(TEST_RATES, threaded=False, extra_modules=[])
    qtbot.addWidget(w)
    w.updates.fetch = lambda url: json.dumps({"tag_name": f"v{__version__}"}).encode()
    w.updates.settings = QSettings(str(tmp_path / "s.ini"), QSettings.Format.IniFormat)
    messages = []
    monkeypatch.setattr(QMessageBox, "information", lambda *a, **k: messages.append(a[2]))
    w.updates.check(manual=True)
    assert messages == [f"У вас последняя версия — {__version__}."]


def test_deleting_rows_removes_exactly_that_row_and_recalculates(window, qtbot):
    """Регрессия: «Удалить» удалял строку под курсором, расчёт оставался с ошибками."""
    from PySide6.QtCore import Qt

    calc, table = window.calculator, window.calculator.form.items_table
    first = table.rows[0]
    first["name"].setText("стеарин")
    first["price"].setText("1000")
    window.calculator.form.add_item()
    window.calculator.form.add_item()
    calc.touched = True
    calc.recalc()
    assert "Позиция 3: укажите наименование" in results(window).status.text()
    table.setCurrentCell(0, 1)  # курсор стоит в заполненной первой строке
    for row in (table.rows[2], table.rows[1]):
        qtbot.mouseClick(row[table.DELETE_KEY], Qt.MouseButton.LeftButton)
    qtbot.waitUntil(lambda: calc.report is not None, timeout=2000)
    assert [r["name"].text() for r in table.rows] == ["стеарин"]
    assert table.rowCount() == 1
    assert not results(window).status.isVisible() or "Позиция" not in results(window).status.text()


def test_deleting_marked_row_does_not_break_recalc(window, qtbot):
    calc, form = window.calculator, window.calculator.form
    form.items_table.rows[0]["name"].setText("Насос")
    form.items_table.rows[0]["price"].setText("100")
    form.add_item()
    calc.touched = True
    calc.recalc()  # пустая вторая строка подсвечена ошибкой
    form.items_table.remove_row(form.items_table.rows[1])
    qtbot.wait(50)  # Qt уничтожает виджеты удалённой строки
    calc.recalc()
    assert calc.report is not None


def test_tnved_code_fills_duty(window):
    from vedcalc.core.models import DutyKind
    from vedcalc.gui.tnved import TnvedSearchDialog

    form = window.calculator.form
    r = form.items_table.rows[0]
    r["name"].setText("Говядина")
    r["price"].setText("4")
    r["quantity"].setText("1000")
    r["unit"].set_value(Unit.KG)
    r["tnved"].setText("0201100008")  # 50%, но не менее 1 EUR за 1 кг
    assert r["duty"].text() == "50" and r["duty"].isReadOnly()
    assert not r["duty_qty"].isEnabled()  # за кг — количество берётся из веса
    window.calculator.recalc()
    assert window.calculator.report.cost.items[0].duty > 0
    # Код со ставкой за пары — нужно «кол-во в ед. ставки»
    r["tnved"].setText("6401100000")
    assert r["duty_qty"].isEnabled() and r["duty_qty"].placeholderText() == "пар"
    # Адвалорная ставка — можно править вручную
    r["tnved"].setText("8507202000")
    assert r["duty"].text() == "3" and not r["duty"].isReadOnly()
    r["duty"].setText("7")
    shipment, _ = form.read()
    assert shipment.items[0].duty.kind is DutyKind.AD_VALOREM
    assert shipment.items[0].duty.percent == 7 and not shipment.items[0].duty.source
    # Неизвестный код подсвечивается, ставка остаётся ручной
    r["tnved"].setText("0000000000")
    assert r["tnved"].unknown and r["tnved"].property("tone") == "warnBorder"
    # Подбор по названию
    dialog = TnvedSearchDialog("аккумуляторы свинцовые")
    assert dialog.results and dialog.selected is not None
    assert "tws.by" in dialog.link.text()
