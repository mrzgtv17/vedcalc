"""Пакетный расчёт WB: чтение XLSX/CSV, колонки, значения из формы, подсветка, выгрузка."""

from dataclasses import replace
from decimal import Decimal as D

import pytest
from openpyxl import Workbook, load_workbook

from vedcalc.modules.wb_unit.batch import (
    BatchFileError,
    build_input,
    evaluate_batch,
    level,
    parse_rows,
    read_table,
    summarize,
    write_results,
    write_template,
)
from vedcalc.modules.wb_unit.calc import Scheme, Target, WbInput, solve

HEADER = ["Название", "Длина, см", "Ширина, см", "Высота, см", "Себестоимость, ₽",
          "Цена до скидки, ₽"]  # fmt: skip


def table(*rows, header=HEADER):
    return [list(header), *[list(map(str, r)) for r in rows]]


def test_template_roundtrip(tmp_path):
    path = tmp_path / "t.xlsx"
    write_template(path)
    rows = parse_rows(read_table(path))
    assert [r.name for r in rows] == ["Кружка керамическая 350 мл", "Плед флисовый 150×200"]
    assert rows[0].values["retail_price"] == D(990) and rows[0].subject == "Кружки"
    assert rows[1].values["buyout_pct"] == D(85) and not rows[0].errors


def test_csv_semicolon_cp1251_and_aliases(tmp_path):
    path = tmp_path / "list.csv"
    text = "Наименование;Длина;Ширина;Высота;Закупка;Цена;Комиссия FBO (%)\n" \
           "Кружка;12;10;11;180,5;990;19,5\n\n"  # fmt: skip
    path.write_bytes(text.encode("cp1251"))
    (row,) = parse_rows(read_table(path))
    assert row.name == "Кружка" and row.values["cost_price"] == D("180.5")
    assert row.values["commission_fbo_pct"] == D("19.5")


def test_csv_comma(tmp_path):
    path = tmp_path / "list.csv"
    path.write_text("name,length,width,height,cost,price\nA,10,10,10,100,500\n", "utf-8")
    (row,) = parse_rows(read_table(path))
    assert row.values["retail_price"] == D(500)


def test_missing_required_columns():
    with pytest.raises(BatchFileError, match="Себестоимость"):
        parse_rows([["Название", "Длина", "Ширина", "Высота"], ["A", "1", "1", "1"]])
    with pytest.raises(BatchFileError):
        parse_rows([[""], []])
    with pytest.raises(BatchFileError, match="нет строк"):
        parse_rows(table())


def test_row_errors_do_not_stop_others():
    rows = parse_rows(table(["A", "10", "abc", "10", "100", "500"],
                            ["B", "10", "10", "10", "", "500"],
                            ["C", "10", "10", "10", "100", "500"]))  # fmt: skip
    assert "«abc» — не число (Ширина, см)" in rows[0].errors
    assert rows[1].errors == ("не заполнено: Себестоимость, ₽",)
    results = evaluate_batch(rows, WbInput())
    assert [r.ok for r in results] == [False, False, True]
    assert summarize(results, D(10)).errors == 2


def test_values_from_file_override_form_and_match_single_calc():
    (row,) = parse_rows(table(["Кружка", "12", "10", "11", "180", "990"],
                              header=[*HEADER, "Скидка, %"]))  # fmt: skip
    base = WbInput(seller_discount_pct=D(25), commission_fbo_pct=D(24))
    (res,) = evaluate_batch([row], base)
    expected = solve(replace(base, name="Кружка", length_cm=D(12), width_cm=D(10),
                             height_cm=D(11), cost_price=D(180), retail_price=D(990)))  # fmt: skip
    assert res.result.profit == expected.profit and res.result.price == expected.price
    assert res.commission_source == "форма"
    # Сравнение схем: та же цена по FBS
    other = solve(replace(expected.inp, scheme=Scheme.FBS))
    assert res.other.profit == other.profit


def test_subject_commission_from_api():
    (row,) = parse_rows(table(["Кружка", "12", "10", "11", "180", "990", "Кружки"],
                              header=[*HEADER, "Предмет WB"]))  # fmt: skip
    inp, source = build_input(row, WbInput(), lambda name: (D(19), D(22)))
    assert (inp.commission_fbo_pct, inp.commission_fbs_pct, source) == (D(19), D(22),
                                                                         "предмет WB")  # fmt: skip
    inp, source = build_input(row, WbInput(), lambda name: None)
    assert inp.commission_fbo_pct == WbInput().commission_fbo_pct and source == "форма"


def test_price_from_target_when_column_empty():
    (row,) = parse_rows(table(["Кружка", "12", "10", "11", "180", ""]))
    (res,) = evaluate_batch([row], WbInput())  # в форме «цена задана» — цены нет
    assert not res.ok and "нет цены" in res.error
    base = WbInput(target=Target.MARGIN, target_value=D(20))
    (res,) = evaluate_batch([row], base)
    assert res.ok and abs(res.result.margin - D("0.2")) < D("0.001")


def test_levels_and_summary():
    rows = parse_rows(table(["Убыток", "35", "30", "8", "1500", "1600"],
                            ["Мало", "12", "10", "11", "180", "650"],
                            ["Хорошо", "12", "10", "11", "180", "990"]))  # fmt: skip
    results = evaluate_batch(rows, WbInput())
    levels = [level(r, D(10)) for r in results]
    assert levels == ["loss", "low", "ok"]
    s = summarize(results, D(10))
    assert s.losses == 1 and s.low_margin == 1 and s.computed == 3
    assert s.profit_sum == sum(r.result.profit for r in results)


def test_write_results(tmp_path):
    rows = parse_rows(table(["A", "12", "10", "11", "180", "990"], ["B", "x", "1", "1", "1", "1"]))
    results = evaluate_batch(rows, WbInput())
    path = tmp_path / "r.xlsx"
    write_results(results, Scheme.FBO, D(10), path)
    ws = load_workbook(path)["Результаты"]
    assert ws.cell(1, 10).value == "Прибыль по FBS"
    assert ws.cell(2, 2).value == "A" and ws.cell(2, 7).value == float(results[0].result.profit)
    assert "не число" in ws.cell(3, 11).value


def test_xlsx_numbers_read_as_text(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.append(HEADER)
    ws.append(["A", 12, 10.5, 11, 180.25, 990])
    path = tmp_path / "n.xlsx"
    wb.save(path)
    (row,) = parse_rows(read_table(path))
    assert row.values["width_cm"] == D("10.5") and row.values["cost_price"] == D("180.25")


# --- вкладка «WB» ----------------------------------------------------------------------


@pytest.fixture
def tab(qtbot):
    pytest.importorskip("PySide6")
    from vedcalc.modules.wb_unit.gui import WbTab

    w = WbTab(threaded=False)
    qtbot.addWidget(w)
    return w


def test_list_mode_in_tab(tab, tmp_path):
    path = tmp_path / "t.xlsx"
    tab.batch.save_template(path)
    assert path.is_file()
    tab.mode_buttons["list"].click()
    assert tab.mode == "list" and tab.pages.currentIndex() == 1
    modified = []
    tab.modified.connect(lambda: modified.append(1))
    assert tab.batch.load_file(path)
    tab.recalc()
    assert tab.batch.table.rowCount() == 2 and modified
    assert "товаров 2" in tab.batch.info.text()
    # Общие параметры — из формы: схема FBS меняет расчёт
    tab.scheme_combo.set_value(Scheme.FBS)
    tab.recalc()
    assert tab.batch.table.horizontalHeaderItem(9).text() == "Прибыль FBO"
    out = tab.batch.export(tmp_path / "out.xlsx")
    assert load_workbook(out)["Результаты"].cell(2, 2).value
    # Ошибка формы — показывается в режиме списка
    tab._fields["buyout_pct"].setText("abc")
    tab.recalc()
    assert tab.batch.table.rowCount() == 0 and "Исправьте форму" in tab.batch.info.text()


def test_list_saved_in_project_state(tab, tmp_path):
    path = tmp_path / "t.xlsx"
    tab.batch.save_template(path)
    tab.mode_buttons["list"].click()
    tab.batch.load_file(path)
    tab.batch.low_margin.setText("15")
    state = tab.save_state()
    assert state["mode"] == "list" and state["batch"]["file"] == "t.xlsx"
    tab.load_state(None)
    assert tab.mode == "single" and tab.batch.rows == []
    tab.load_state(state)
    assert tab.mode == "list" and len(tab.batch.rows) == 2
    assert tab.batch.low_margin.text() == "15" and tab.batch.table.rowCount() == 2


def test_bad_file_message(tab, tmp_path):
    path = tmp_path / "bad.csv"
    path.write_text("a;b\n1;2\n", "utf-8")
    assert not tab.batch.load_file(path)
    assert "Нет обязательных колонок" in tab.batch.info.text()
