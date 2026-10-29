"""Ступенчатый прайс: цена и прибыль за единицу по порогам, проверка ввода, файл, экран."""

from decimal import Decimal as D

import pytest
from conftest import TEST_RATES, control_shipment
from test_costing import two_items
from test_whatif import report_for

from vedcalc.core.io import shipment_from_dict, shipment_to_dict
from vedcalc.core.models import PriceTier, TierPricing
from vedcalc.core.money import round_money
from vedcalc.core.tiers import tier_label, tier_table, validate
from vedcalc.core.validation import InputError


def test_tier_without_discount_and_order_cost_matches_scenario():
    report = report_for(two_items())
    view = report.unit_views()[0]
    table = tier_table(report, TierPricing([PriceTier(D(1), D(0))]))
    row = table.rows[0]
    assert table.base_price_net == view.price.net
    assert row.price_net == view.price.net and row.vat == view.price.vat
    assert row.net_profit == view.price.net_profit
    assert tier_label(row, table.unit_label) == f"от 1 {view.unit_label}"


def test_discount_and_order_cost_per_unit():
    report = report_for(control_shipment())
    view = report.unit_views()[0]
    params = TierPricing([PriceTier(D(10), D(0)), PriceTier(D(100), D(5))], D(3000))
    small, big = tier_table(report, params).rows
    base = view.price.net
    assert small.order_cost == D(300) and big.order_cost == D(30)
    assert big.price_net == round_money(base * D("0.95"))
    gross = small.price_net - view.cost_net - view.financing - small.order_cost
    assert small.gross_profit == gross
    assert small.net_profit == gross - round_money(gross * D("0.25"))
    assert small.margin == gross / small.price_net
    assert small.price_gross == small.price_net + round_money(small.price_net * D("0.22"))
    assert small.over_batch and big.over_batch  # в контрольном примере 1 шт


def test_loss_tier_is_flagged():
    report = report_for(control_shipment())
    table = tier_table(report, TierPricing([PriceTier(D(1), D(50))]))
    assert table.rows[0].loss and table.rows[0].net_profit < 0
    assert any("убытк" in n for n in table.notes)


def test_tier_item_selects_position():
    report = report_for(two_items())
    views = report.unit_views()
    table = tier_table(report, TierPricing([PriceTier(D(1))], item=1))
    assert table.item_name == views[1].name


def test_validation():
    bad = TierPricing([PriceTier(D(0)), PriceTier(D(50), D(100)), PriceTier(D(20))], D(-1))
    messages = [i.message for i in validate(bad)]
    assert any("больше 0" in m for m in messages)
    assert any("по возрастанию" in m for m in messages)
    assert any("от 0 до 99,99" in m for m in messages)
    assert any("Расходы на заказ" in m for m in messages)
    with pytest.raises(InputError):
        tier_table(report_for(two_items()), bad)
    assert validate(TierPricing()) == []


def test_tiers_roundtrip_in_file():
    s = control_shipment()
    s.tier_pricing = TierPricing([PriceTier(D(5), D("2.5"))], D(1500), 0)
    back = shipment_from_dict(shipment_to_dict(s))
    assert back.tier_pricing == s.tier_pricing
    data = shipment_to_dict(s)
    del data["shipment"]["tier_pricing"]
    assert shipment_from_dict(data).tier_pricing == TierPricing()  # файлы 0.12 и раньше


# --- экран -------------------------------------------------------------------------------


@pytest.fixture
def window(qtbot, monkeypatch):
    pytest.importorskip("PySide6")
    from PySide6.QtWidgets import QMessageBox

    from vedcalc.gui.main_window import MainWindow

    monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: QMessageBox.StandardButton.Discard)
    w = MainWindow(TEST_RATES, threaded=False, extra_modules=[])
    qtbot.addWidget(w)
    w.open_example()
    return w


def test_tier_card_on_screen(window):
    view = window.pricing.results.tier_view
    assert view.table is not None and len(view.table.rows) == 3
    window.pricing.results.setCurrentIndex(0)
    assert not view.isHidden()
    assert view.findChild(object, "tier:2:1").text() == "3,0%"


def test_editing_tiers_saves_state_without_breaking_price(window):
    view = window.pricing.results.tier_view
    view.order_cost.setText("3000")
    assert view.params.order_cost == D(3000)
    assert window.calculator.form.pricing.tiers.order_cost == D(3000)
    assert window.dirty
    view.rows[1][0].setText("5")  # 5 < 10 — не по возрастанию
    assert view.rows[1][0].property("tone") == "fieldError"
    assert view.params.tiers[1].min_qty == D(50)  # остаётся последний корректный ввод
    assert window.pricing.report.target is not None
    view.rows[1][0].setText("60")
    assert view.params.tiers[1].min_qty == D(60)
    view.add_tier()
    assert len(view.params.tiers) == 4 and view.params.tiers[3].min_qty == D(200)
    view._remove(view.rows[0][0])
    assert [t.min_qty for t in view.params.tiers] == [D(60), D(100), D(200)]


def test_tiers_survive_save_and_open(window, tmp_path):
    view = window.pricing.results.tier_view
    view.order_cost.setText("1200")
    path = tmp_path / "tiers.vedcalc.json"
    window._save_to(path)
    window.open_example()
    assert view.params.order_cost == 0
    window.open_file(path)
    assert view.params.order_cost == D(1200)
    assert view.order_cost.value() == D(1200)


# --- экспорт и КП ------------------------------------------------------------------------


def _table():
    report = report_for(control_shipment())
    return report, tier_table(report, TierPricing(order_cost=D(3000)))


def test_price_report_html_and_xlsx_include_tiers(tmp_path):
    from openpyxl import load_workbook

    from vedcalc.reports.html import render_price_html
    from vedcalc.reports.xlsx import write_price_xlsx

    report, table = _table()
    html = render_price_html(report, tiers=table)
    assert "Ступенчатый прайс" in html and "от 50 шт" in html and "3,0%" in html
    assert "Ступенчатый прайс" not in render_price_html(report)
    path = tmp_path / "price.xlsx"
    write_price_xlsx(report, path, tiers=table)
    ws = load_workbook(path)["Прайс по объёму"]
    labels = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1)]
    assert "от 100 шт" in labels


def test_offer_shows_volume_prices_without_profit(tmp_path):
    from vedcalc.core.offer import OfferInput, Requisites, build_offer
    from vedcalc.reports.offer_docx import write_offer_docx
    from vedcalc.reports.offer_html import render_offer_html

    report, table = _table()
    params = TierPricing(order_cost=D(3000))
    hidden = build_offer(report, OfferInput(), Requisites(), tier_pricing=params)
    assert hidden.tiers == ()
    offer = build_offer(report, OfferInput(show_tiers=True), Requisites(), tier_pricing=params)
    assert [t.label for t in offer.tiers] == ["от 10 шт", "от 50 шт", "от 100 шт"]
    assert [t.discount for t in offer.tiers] == ["—", "3%", "5%"]
    assert [t.unit_price for t in offer.tiers] == [r.price_gross for r in table.rows]
    html = render_offer_html(offer)
    assert offer.tiers_title in html
    assert "прибыл" not in html.lower()
    write_offer_docx(offer, tmp_path / "kp.docx")
    from docx import Document

    text = "\n".join(c.text for t in Document(tmp_path / "kp.docx").tables for c in t._cells)
    assert "от 100 шт" in text


def test_offer_dialog_tiers_checkbox(qtbot, tmp_path):
    pytest.importorskip("PySide6")
    from PySide6.QtCore import QSettings

    from vedcalc.gui.offer import OfferDialog

    report, _ = _table()
    settings = QSettings(str(tmp_path / "s.ini"), QSettings.Format.IniFormat)
    plain = OfferDialog(report, directory=tmp_path, settings=settings)
    qtbot.addWidget(plain)
    assert not plain.show_tiers.isEnabled()
    dialog = OfferDialog(report, directory=tmp_path, settings=settings,
                         tier_pricing=TierPricing())  # fmt: skip
    qtbot.addWidget(dialog)
    dialog.show_tiers.setChecked(True)
    dialog.refresh()
    assert len(dialog.offer.tiers) == 3
