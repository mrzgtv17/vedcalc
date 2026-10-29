"""История вариантов и сравнение: снимки в файле проекта, таблица по статьям, окно."""

import copy
from decimal import Decimal as D

import pytest
from conftest import TEST_RATES, control_shipment
from test_costing import two_items

from vedcalc.core.compare import Better, Kind, compare, evaluate
from vedcalc.core.fees import load_fee_scale
from vedcalc.core.io import project_from_dict, project_to_dict, snapshots_from_dict
from vedcalc.core.models import PriceMode, PricingSource, Snapshot
from vedcalc.core.money import Money


@pytest.fixture(scope="module")
def scale():
    return load_fee_scale()


def rail(shipment):
    s = copy.deepcopy(shipment)
    s.name = "Ж/д"
    s.costs[0].amount = Money.rub(45000)
    return s


def rows(c):
    return {r.label: r for r in c.rows}


def test_comparison_by_articles_and_deltas(scale):
    sea = control_shipment()
    c = compare([evaluate("Море", sea, TEST_RATES, scale),
                 evaluate("Ж/д", rail(sea), TEST_RATES, scale)])  # fmt: skip
    r = rows(c)
    assert c.names == ("Море", "Ж/д") and not c.errors
    assert r["Расходы до границы (в таможенной стоимости)"].delta(1) == D(15000)
    cost = r["Себестоимость партии без НДС"]
    assert cost.values[0] == D("147731.00") and cost.better is Better.LOW and cost.strong
    assert cost.verdict(1) == -1  # дороже — хуже
    assert r["Пошлина"].delta(1) == D(750)
    assert r["Стоимость товара"].verdict(1) == 0
    assert "Акциз" not in r  # везде 0 — строку не показываем
    assert r["Маржинальность"].kind is Kind.PCT
    assert r["Себестоимость 1 шт без НДС"].values[1] == D("163481.00")
    assert cost.delta(0) is None


def test_profit_is_better_when_higher_at_same_price(scale):
    sea = control_shipment()
    sea.sale.target_mode, sea.sale.target_value = PriceMode.PRICE, D(200000)
    c = compare([evaluate("Море", sea, TEST_RATES, scale),
                 evaluate("Ж/д", rail(sea), TEST_RATES, scale)])  # fmt: skip
    profit = rows(c)["Чистая прибыль партии"]
    assert profit.better is Better.HIGH and profit.delta(1) < 0 and profit.verdict(1) == -1
    assert not any("от себестоимости" in n for n in c.notes)


def test_notes_and_errors(scale):
    a = control_shipment()
    b = rail(a)
    b.calc_date = a.calc_date.replace(day=a.calc_date.day + 1)
    b.sale.target_mode, b.sale.target_value = PriceMode.PRICE, D(200000)
    bad = copy.deepcopy(a)
    bad.items[0].quantity = D(0)
    c = compare([evaluate("A", a, TEST_RATES, scale), evaluate("B", b, TEST_RATES, scale),
                 evaluate("Ошибка", bad, TEST_RATES, scale)])  # fmt: skip
    assert any("Даты расчёта разные" in n for n in c.notes)
    assert any("Сценарии цены разные" in n for n in c.notes)
    assert c.errors and c.errors[0][0] == "Ошибка"
    assert rows(c)["Стоимость товара"].values[2] is None


def test_manual_cost_variant(scale):
    manual = control_shipment()
    manual.pricing_source = PricingSource.MANUAL
    manual.manual_cost.quantity = D(1)
    manual.manual_cost.unit_cost_net = D(150000)
    c = compare([evaluate("Импорт", control_shipment(), TEST_RATES, scale),
                 evaluate("Вручную", manual, TEST_RATES, scale)])  # fmt: skip
    r = rows(c)
    assert r["Пошлина"].values[1] is None
    assert r["Себестоимость партии без НДС"].values[1] == D(150000)


def test_multi_item_variants_skip_unit_rows(scale):
    c = compare([evaluate("1", two_items(), TEST_RATES, scale),
                 evaluate("2", two_items(), TEST_RATES, scale)])  # fmt: skip
    assert not any(label.startswith("Себестоимость 1 ") for label in rows(c))


def test_snapshots_roundtrip_in_project_file():
    s = control_shipment()
    snaps = [Snapshot("Море", copy.deepcopy(s)), Snapshot("Ж/д", rail(s))]
    data = project_to_dict(s, None, snaps)
    back = snapshots_from_dict(data)
    assert [x.name for x in back] == ["Море", "Ж/д"]
    assert back[1].shipment.costs[0].amount == Money.rub(45000)
    assert back[0].created == snaps[0].created
    project_from_dict(data)  # основной расчёт читается как прежде
    assert snapshots_from_dict(project_to_dict(s)) == []
    assert "snapshots" not in project_to_dict(s)


# --- экран -------------------------------------------------------------------------------


@pytest.fixture
def window(qtbot, monkeypatch):
    pytest.importorskip("PySide6")
    from PySide6.QtWidgets import QMessageBox

    from vedcalc.gui.main_window import MainWindow

    monkeypatch.setattr(QMessageBox, "question", lambda *a, **k: QMessageBox.StandardButton.Discard)
    w = MainWindow(TEST_RATES, threaded=False, extra_modules=[])
    qtbot.addWidget(w)
    w.open_example()
    return w


def test_remember_compare_and_open_variant(window, qtbot, tmp_path):
    from vedcalc.gui.compare import CURRENT, CompareDialog

    window.remember_variant("Море")
    assert window.dirty and window.snapshots[0].name == "Море"
    window.calculator.form.costs_table.rows[0]["amount"].setText("45000")
    window.calculator.recalc()
    window.remember_variant("Ж/д")

    dialog = CompareDialog(window)
    qtbot.addWidget(dialog)
    assert dialog.comparison.names == ("Море", "Ж/д")  # два варианта — без текущего
    path = dialog.export(tmp_path / "cmp.xlsx")
    from openpyxl import load_workbook

    ws = load_workbook(path)["Сравнение"]
    assert "Ж/д − Море" in {c.value for row in ws.iter_rows() for c in row}

    dialog.list.item(0).setCheckState(dialog.list.item(0).checkState().Checked)
    assert dialog.comparison.names == (CURRENT, "Море", "Ж/д")

    dialog.list.setCurrentRow(2)
    dialog.rename("Железная дорога")
    assert window.snapshots[1].name == "Железная дорога"

    dialog.list.setCurrentRow(1)
    dialog.open_in_editor()
    assert window.calculator.form.costs_table.rows[0]["amount"].text() == "30000"
    assert len(window.snapshots) == 2


def test_variants_saved_with_project_and_added_from_file(window, qtbot, tmp_path):
    from vedcalc.gui.compare import CompareDialog

    window.remember_variant("Море")
    path = tmp_path / "p.vedcalc.json"
    window._save_to(path)
    window.open_example()
    assert window.snapshots == []
    window.open_file(path)
    assert [s.name for s in window.snapshots] == ["Море"]

    dialog = CompareDialog(window)
    qtbot.addWidget(dialog)
    added = dialog.add_from_file(path)
    assert added.name == "Контрольный пример" and len(window.snapshots) == 2
    dialog.list.setCurrentRow(2)
    dialog.delete()
    assert [s.name for s in window.snapshots] == ["Море"]
