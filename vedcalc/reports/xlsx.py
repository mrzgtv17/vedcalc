"""Экспорт отчёта в XLSX (openpyxl). Суммы пишутся числами, а не текстом."""

from __future__ import annotations

from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.worksheet import Worksheet

from vedcalc.core.explain import fmt_money
from vedcalc.core.report import CostReport, PriceReport, Row
from vedcalc.core.tiers import TierTable, tier_label
from vedcalc.core.whatif import WhatIf

MONEY = "#,##0.00"
HEADER_FILL = PatternFill("solid", fgColor="EEF2F7")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=14)
SECTION = Font(bold=True, size=12, color="0B5394")


def _write_rows(ws: Worksheet, start: int, title: str, rows: tuple[Row, ...]) -> int:
    ws.cell(start, 1, title).font = SECTION
    r = start + 1
    for row in rows:
        ws.cell(r, 1, row.label)
        cell = ws.cell(r, 2, row.value if isinstance(row.value, Decimal) else row.text)
        if isinstance(row.value, Decimal):
            cell.number_format = MONEY
        else:
            cell.alignment = Alignment(horizontal="right")
        if row.hint:
            ws.cell(r, 3, row.hint).font = Font(color="6E6E73")
        if row.strong:
            ws.cell(r, 1).font = BOLD
            cell.font = BOLD
        r += 1
    return r + 1


def _lines(ws: Worksheet, r: int, heading: str, lines) -> int:
    for c, h in enumerate((heading, "Сумма, руб.", "Формула"), 1):
        cell = ws.cell(r, c, h)
        cell.font = BOLD
        cell.fill = HEADER_FILL
    r += 1
    for line in lines:
        ws.cell(r, 1, line.label)
        value = ws.cell(r, 2, line.value)
        value.number_format = MONEY
        ws.cell(r, 3, line.formula).font = Font(color="76746D")
        if line.strong:
            ws.cell(r, 1).font = BOLD
            value.font = BOLD
        r += 1
    return r + 1


def _header(ws: Worksheet, title: str, subtitle: str, warnings) -> int:
    ws.cell(1, 1, title).font = TITLE
    ws.cell(2, 1, subtitle)
    r = 4
    for w in warnings:
        ws.cell(r, 1, f"⚠ {w}").font = Font(color="9C4A00")
        r += 1
    return r + (1 if warnings else 0)


def _widths(ws: Worksheet, *widths: int) -> None:
    for letter, width in zip("ABCD", widths, strict=False):
        ws.column_dimensions[letter].width = width


def _steps_sheet(wb: Workbook, groups) -> None:
    steps = wb.create_sheet("Детализация")
    _widths(steps, 44, 110)
    r = 1
    for title, group in groups:
        steps.cell(r, 1, title).font = SECTION
        r += 1
        for st in group:
            steps.cell(r, 1, st.title)
            steps.cell(r, 2, st.formula)
            r += 1
        r += 1


def write_cost_xlsx(report: CostReport, path: Path) -> None:
    """Себестоимость импорта: «За единицу», «Итоги», «Позиции», «Детализация»."""
    from vedcalc.reports.html import load_warning

    wb = Workbook()
    s = report.shipment
    warnings = [*report.cost.warnings, *filter(None, [load_warning(report)])]
    subtitle = (
        f"Себестоимость импорта · дата расчёта {s.calc_date:%d.%m.%Y} · {s.delivery_mode.label}"
    )

    ws = wb.active
    ws.title = "За единицу"
    r = _header(ws, s.name, subtitle, warnings)
    for view in report.unit_views():
        ws.cell(r, 1, f"{view.name} — {view.per}").font = SECTION
        r = _lines(ws, r + 1, "Себестоимость " + view.per, view.cost_lines) + 1
    _widths(ws, 40, 22, 62)

    ws = wb.create_sheet("Итоги")
    r = _header(ws, s.name, subtitle, ())
    r = _write_rows(ws, r, *_section(report.batch_section()))
    if report.load is not None:
        r = _write_rows(ws, r, f"Загрузка контейнера {report.load.spec.title}", report.load_rows())
    _write_rows(ws, r, "Курсы и методика", report.rates_rows() + report.method_rows())
    _widths(ws, 42, 44, 60)

    items = wb.create_sheet("Позиции")
    head = [
        "Позиция", "Количество", "Ед.", "Товар, руб.", "Расходы в ТС", "Прочие расходы",
        "Таможенная стоимость", "Таможенный сбор", "Пошлина", "Акциз", "НДС при ввозе",
        "НДС по услугам", "Себестоимость без НДС", "НДС к вычету",
        "Себестоимость ед. без НДС", "Себестоимость ед. с НДС",
    ]  # fmt: skip
    for c, h in enumerate(head, 1):
        cell = items.cell(1, c, h)
        cell.font = BOLD
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        items.column_dimensions[cell.column_letter].width = 16
    items.column_dimensions["A"].width = 30
    for n, (i, v) in enumerate(zip(report.cost.items, report.unit_views(), strict=True), 2):
        values = [
            i.name, i.quantity, i.unit.label, i.goods, i.costs_in_cv, i.costs_other,
            i.customs_value, i.customs_fee, i.duty, i.excise, i.import_vat, i.services_vat,
            i.cost_net, i.deductible_vat, v.cost_net, v.cost_gross,
        ]  # fmt: skip
        for c, val in enumerate(values, 1):
            cell = items.cell(n, c, val)
            if c >= 4:
                cell.number_format = MONEY
    items.freeze_panes = "B2"

    _steps_sheet(wb, report.step_groups())
    wb.save(path)


def _whatif_sheet(wb: Workbook, report: PriceReport, whatif: WhatIf) -> None:
    ws = wb.create_sheet("Что если")
    r = _header(ws, report.title, f"Цена продажи не меняется: {fmt_money(whatif.price_net)} ₽ "
                f"без НДС за партию. {whatif.scope_text}", ())  # fmt: skip
    for table in whatif.tables:
        ws.cell(r, 1, table.title).font = SECTION
        r += 1
        for c, h in enumerate(whatif.heads, 1):
            ws.cell(r, c, h + (", %" if c == 4 else "")).font = BOLD
        r += 1
        for row in table.rows:
            cost = row.unit_cost if row.unit_cost is not None else row.cost_net
            margin = None if row.margin is None else round(row.margin * 100, 2)
            values = (row.scenario.label, cost, row.net_profit, margin, row.profit_delta)
            for c, value in enumerate(values, 1):
                cell = ws.cell(r, c, value)
                if isinstance(value, Decimal):
                    cell.number_format = "#,##0.00"
                if row.scenario.is_base:
                    cell.font = BOLD
            r += 1
        ws.cell(r, 1, table.note).font = Font(italic=True, color="76746D")
        r += 2
    for note in whatif.notes:
        ws.cell(r, 1, note)
        r += 1
    _widths(ws, 22, 24, 20, 12, 16)


def _tiers_sheet(wb: Workbook, report: PriceReport, t: TierTable) -> None:
    ws = wb.create_sheet("Прайс по объёму")
    subtitle = (
        f"Ступенчатый прайс: {t.item_name}. Базовая цена — по сценарию: "
        f"{fmt_money(t.base_price_net)} ₽ без НДС за 1 {t.unit_label}."
    )
    r = _header(ws, report.title, subtitle, ())
    heads = ["Заказ", "Скидка, %", "Цена без НДС"] + (["Цена с НДС"] if t.vat_payer else [])
    heads += ["Расходы на заказ на ед.", "Прибыль на ед.", "Маржа, %"]
    for c, h in enumerate(heads, 1):
        ws.cell(r, c, h).font = BOLD
    r += 1
    for row in t.rows:
        values = [tier_label(row, t.unit_label), row.discount_pct, row.price_net]
        if t.vat_payer:
            values.append(row.price_gross)
        margin = None if row.margin is None else round(row.margin * 100, 2)
        values += [row.order_cost, row.net_profit, margin]
        for c, value in enumerate(values, 1):
            cell = ws.cell(r, c, value)
            if isinstance(value, Decimal) and c > 2:
                cell.number_format = "#,##0.00"
        r += 1
    r += 1
    for note in t.notes:
        ws.cell(r, 1, note)
        r += 1
    _widths(ws, 18, 12, 16, 16, 20, 16, 12)


def write_price_xlsx(
    report: PriceReport,
    path: Path,
    whatif: WhatIf | None = None,
    tiers: TierTable | None = None,
) -> None:
    """Цена продажи: «За единицу», «Итоги» (+ стоимость денег), «Прайс по объёму»,
    «Что если», «Детализация»."""
    wb = Workbook()
    subtitle = f"Цена продажи · {report.target_title} · {report.basis.source_label}"
    warnings = [report.target_error] if report.target is None and report.target_error else []

    ws = wb.active
    ws.title = "За единицу"
    r = _header(ws, report.title, subtitle, warnings)
    for view in report.unit_views():
        ws.cell(r, 1, f"{view.name} — {view.per}").font = SECTION
        r = _lines(ws, r + 1, "Цена и прибыль " + view.per, view.price.lines) + 1
    _widths(ws, 40, 22, 62)

    ws = wb.create_sheet("Итоги")
    r = _header(ws, report.title, subtitle, ())
    r = _write_rows(ws, r, *_section(report.batch_section()))
    if fin := report.financing_rows():
        r = _write_rows(ws, r, "Стоимость денег", fin)
        periods = report.period_rows()
        if periods:
            for c, h in enumerate(("Дни", "Длительность", "Вложено", "Проценты"), 1):
                ws.cell(r, c, h).font = BOLD
            r += 1
            for row in periods:
                for c, v in enumerate(row, 1):
                    ws.cell(r, c, v)
                r += 1
            r += 1
    _write_rows(ws, r, "Методика", report.method_rows())
    _widths(ws, 42, 30, 60, 18)

    if tiers is not None:
        _tiers_sheet(wb, report, tiers)
    if whatif is not None and whatif.tables:
        _whatif_sheet(wb, report, whatif)
    _steps_sheet(wb, report.step_groups())
    wb.save(path)


def _section(section) -> tuple[str, tuple[Row, ...]]:
    return section.title, section.rows
