"""КП в DOCX (python-docx): то же содержание и те же три стиля, что и в PDF."""

from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from vedcalc.core.explain import fmt_money
from vedcalc.core.offer import Offer, OfferStyle
from vedcalc.reports.offer_html import PALETTES

# Шрифты Word: системный UI-шрифт macOS в docx не встроить — берём близкие стандартные
FONTS = {
    OfferStyle.PROGRAM: "Helvetica Neue",
    OfferStyle.NEUTRAL: "Arial",
    OfferStyle.MONO: "Times New Roman",
}


def _rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.lstrip("#").upper())


def _shade(cell, hex_color: str) -> None:
    props = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color.lstrip("#").upper())
    props.append(shd)


def _borders(table, *, show: bool) -> None:
    """Сетка таблицы: тонкие чёрные линии (монохром) или без линий."""
    props = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:val"), "single" if show else "nil")
        if show:
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), "000000")
        borders.append(el)
    props.append(borders)


def _fixed(table, widths_mm: list[float]) -> None:
    """Фиксированные ширины колонок: без этого Word и Pages ужимают таблицу по содержимому."""
    table.autofit = False
    props = table._tbl.tblPr
    layout = OxmlElement("w:tblLayout")
    layout.set(qn("w:type"), "fixed")
    props.append(layout)
    total = OxmlElement("w:tblW")
    total.set(qn("w:w"), str(Mm(sum(widths_mm)).twips))
    total.set(qn("w:type"), "dxa")
    for old in props.findall(qn("w:tblW")):
        props.remove(old)
    props.append(total)
    grid = table._tbl.tblGrid
    for col, width in zip(grid.findall(qn("w:gridCol")), widths_mm, strict=True):
        col.set(qn("w:w"), str(Mm(width).twips))
    for row in table.rows:
        k = 0
        for tc in row._tr.tc_lst:
            span = tc.grid_span
            tc.width = Mm(sum(widths_mm[k : k + span]))
            k += span


def _bottom_rule(paragraph, hex_color: str, size: int = 8) -> None:
    props = paragraph._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), hex_color.lstrip("#").upper())
    bdr.append(bottom)
    props.append(bdr)


def _text(cell_or_par, text: str, *, bold=False, color=None, size=None, align=None):
    par = cell_or_par.paragraphs[0] if hasattr(cell_or_par, "paragraphs") else cell_or_par
    run = par.add_run(text)
    run.bold = bold
    if color:
        run.font.color.rgb = _rgb(color)
    if size:
        run.font.size = Pt(size)
    if align is not None:
        par.alignment = align
    return run


def write_offer_docx(offer: Offer, path: Path) -> None:
    inp, req = offer.inp, offer.requisites
    p = PALETTES[inp.style]
    mono = inp.style is OfferStyle.MONO

    doc = Document()
    section = doc.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.top_margin, section.bottom_margin = Mm(15), Mm(15)
    section.left_margin, section.right_margin = Mm(18), Mm(15)
    normal = doc.styles["Normal"]
    normal.font.name = FONTS[inp.style]
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONTS[inp.style])
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = _rgb(p.text)
    normal.paragraph_format.space_after = Pt(2)

    # Шапка: логотип слева, компания и реквизиты справа
    logo = req.image("logo") if inp.use_logo else None
    head = doc.add_table(rows=1, cols=2 if logo else 1)
    _borders(head, show=False)
    info = head.rows[0].cells[-1]
    if logo:
        head.rows[0].cells[0].paragraphs[0].add_run().add_picture(str(logo), width=Mm(40))
    align = WD_ALIGN_PARAGRAPH.RIGHT if logo else WD_ALIGN_PARAGRAPH.LEFT
    _text(info, req.company or "Ваша компания", bold=True, size=13, align=align)
    for line in (req.tax_ids, req.address, req.contacts):
        if line:
            _text(info.add_paragraph(), line, color=p.muted, size=8.5, align=align)
    _fixed(head, [60, 117] if logo else [177])
    _bottom_rule(doc.add_paragraph(), p.rule)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(8)
    _text(title, inp.subject or "Коммерческое предложение", bold=True, color=p.accent, size=18)
    sub = doc.add_paragraph()
    _text(sub, f"от {inp.offer_date:%d.%m.%Y}", color=p.muted)
    to = ", ".join(x for x in (inp.client_company, inp.client_person) if x.strip())
    if to:
        _text(sub, " · для: ", color=p.muted)
        _text(sub, to, bold=True, color=p.muted)
    doc.add_paragraph("Предлагаем к поставке следующие товары:")

    heads = ["№", "Наименование"]
    if inp.show_tnved:
        heads.append("Код ТН ВЭД")
    heads += ["Кол-во", offer.price_header, "Сумма, ₽"]
    num_cols = {len(heads) - 3, len(heads) - 2, len(heads) - 1}
    table = doc.add_table(rows=1, cols=len(heads))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    _borders(table, show=mono)
    for k, h in enumerate(heads):
        cell = table.rows[0].cells[k]
        _text(cell, h, bold=True, color=p.head_text, size=9,
              align=WD_ALIGN_PARAGRAPH.RIGHT if k in num_cols else None)  # fmt: skip
        if not mono:
            _shade(cell, p.head_bg)
    for n, line in enumerate(offer.lines, 1):
        values = [str(n), line.name]
        if inp.show_tnved:
            values.append(line.tnved or "—")
        values += [line.quantity_text, fmt_money(line.unit_price), fmt_money(line.amount)]
        cells = table.add_row().cells
        for k, value in enumerate(values):
            _text(cells[k], value, align=WD_ALIGN_PARAGRAPH.RIGHT if k in num_cols else None)
            if n % 2 == 0 and not mono:
                _shade(cells[k], p.zebra)
    span = len(heads) - 1

    def summary_row(label: str, value: str, *, bold: bool, shade: str | None) -> None:
        cells = table.add_row().cells
        merged = cells[0].merge(cells[span - 1])
        _text(merged, label, bold=bold)
        _text(cells[span], value, bold=bold, align=WD_ALIGN_PARAGRAPH.RIGHT)
        if shade and not mono:
            _shade(merged, shade)
            _shade(cells[span], shade)

    summary_row("Итого", fmt_money(offer.total), bold=True, shade=p.total_bg)
    if offer.vat_payer:
        rate = f"{offer.vat_rate.normalize():f}%"
        summary_row(f"В том числе НДС {rate}", fmt_money(offer.vat_total), bold=False, shade=None)
    else:
        cells = table.add_row().cells
        _text(cells[0].merge(cells[span]), offer.vat_text, color=p.muted)
    name_mm = 64 if inp.show_tnved else 88
    _fixed(table, [9, name_mm] + ([24] if inp.show_tnved else []) + [18, 34, 28])

    if offer.tiers:
        doc.add_paragraph()
        _text(doc.add_paragraph(), offer.tiers_title, bold=True)
        tt = doc.add_table(rows=1, cols=3)
        _borders(tt, show=mono)
        for k, h in enumerate(("Объём заказа", "Скидка", offer.price_header)):
            cell = tt.rows[0].cells[k]
            _text(cell, h, bold=True, color=p.head_text, size=9,
                  align=WD_ALIGN_PARAGRAPH.RIGHT if k else None)  # fmt: skip
            if not mono:
                _shade(cell, p.head_bg)
        for n, t in enumerate(offer.tiers, 1):
            cells = tt.add_row().cells
            for k, value in enumerate((t.label, t.discount, fmt_money(t.unit_price))):
                _text(cells[k], value, align=WD_ALIGN_PARAGRAPH.RIGHT if k else None)
                if n % 2 == 0 and not mono:
                    _shade(cells[k], p.zebra)
        _fixed(tt, [50, 30, 44])

    terms = offer.terms()
    if terms:
        doc.add_paragraph()
        t = doc.add_table(rows=0, cols=2)
        _borders(t, show=False)
        for label, value in terms:
            cells = t.add_row().cells
            _text(cells[0], label, color=p.muted)
            _text(cells[1], value)
        _fixed(t, [60, 117])
    if inp.comment.strip():
        doc.add_paragraph()
        doc.add_paragraph(inp.comment.strip())

    if req.signer_name or req.signer_position:
        doc.add_paragraph()
        sign = doc.add_table(rows=1, cols=3)
        _borders(sign, show=False)
        left, middle, right = sign.rows[0].cells
        _fixed(sign, [65, 67, 45])
        _text(left, req.signer_position)
        _text(left.add_paragraph(), req.company)
        par = middle.paragraphs[0]
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
        stamp = req.image("stamp") if inp.use_stamp else None
        signature = req.image("signature") if inp.use_signature else None
        if stamp:
            par.add_run().add_picture(str(stamp), width=Mm(32))
        if signature:
            par.add_run().add_picture(str(signature), width=Mm(30))
        _text(right, req.signer_name, align=WD_ALIGN_PARAGRAPH.RIGHT)
        for cell in (left, middle, right):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.BOTTOM
    if req.bank:
        doc.add_paragraph()
        _text(doc.add_paragraph(), f"Банковские реквизиты: {req.bank}", color=p.muted, size=8.5)
    doc.save(str(path))
