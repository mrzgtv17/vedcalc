"""Пакетный расчёт WB: список товаров из XLSX/CSV → прибыль и маржа по каждому.

Общие параметры (схема, налог, тарифы, склад, выкуп…) берутся из формы вкладки «WB»;
в файле — то, что отличается у товаров. Пустая ячейка — значение из формы.
"""

from __future__ import annotations

import csv
import io
import re
from collections.abc import Callable
from dataclasses import dataclass, field, replace
from decimal import Decimal, InvalidOperation
from pathlib import Path

from vedcalc.modules.wb_unit.calc import (
    Scheme,
    Target,
    Tariffs,
    WbInput,
    WbInputError,
    WbResult,
    solve,
)

DEFAULT_LOW_MARGIN = Decimal(10)  # маржа ниже — жёлтая подсветка, %


@dataclass(frozen=True, slots=True)
class Field:
    key: str
    header: str  # заголовок в шаблоне
    aliases: tuple[str, ...]
    required: bool = False
    hint: str = ""


FIELDS = (
    Field("name", "Название", ("название", "наименование", "товар", "name", "артикул"), True),
    Field("length_cm", "Длина, см", ("длина", "length", "д"), True, "Упаковка"),
    Field("width_cm", "Ширина, см", ("ширина", "width", "ш"), True, "Упаковка"),
    Field("height_cm", "Высота, см", ("высота", "height", "в"), True, "Упаковка"),
    Field("cost_price", "Себестоимость, ₽", ("себестоимость", "закупка", "cost"), True,
          "Закупка и доставка до склада"),
    Field("retail_price", "Цена до скидки, ₽", ("цена до скидки", "цена", "price"), False,
          "Пусто — цена подбирается под цель из формы (маржа или прибыль)"),
    Field("seller_discount_pct", "Скидка, %", ("скидка", "скидка продавца", "discount")),
    Field("subject", "Предмет WB", ("предмет", "предмет wb", "категория", "subject"), False,
          "Комиссия — из тарифов WB по API, если колонки комиссий пустые"),
    Field("commission_fbo_pct", "Комиссия FBO, %", ("комиссия fbo", "fbo")),
    Field("commission_fbs_pct", "Комиссия FBS, %", ("комиссия fbs", "fbs")),
    Field("buyout_pct", "Выкуп, %", ("выкуп", "buyout")),
    Field("ads_pct", "Реклама, %", ("реклама", "ads")),
    Field("other_per_unit", "Прочие, ₽", ("прочие", "прочие расходы", "other")),
)  # fmt: skip
DECIMAL_KEYS = tuple(f.key for f in FIELDS if f.key not in ("name", "subject"))


class BatchFileError(ValueError):
    pass


def _norm(text: str) -> str:
    """«Длина, см» → «длина»; «Комиссия FBO (%)» → «комиссия fbo»."""
    text = str(text or "").lower().replace("ё", "е").replace("\xa0", " ")
    text = re.sub(r"\(.*?\)", "", text)
    text = re.sub(r",\s*(%|₽|руб\.?|см|cm)\s*$", "", text.strip())
    return " ".join(text.replace("%", "").replace("₽", "").split())


_ALIASES = {_norm(a): f.key for f in FIELDS for a in (f.header, *f.aliases)}


def read_table(path: Path) -> list[list[str]]:
    """Строки файла как текст: XLSX — первый лист, CSV — разделитель ; , или табуляция."""
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        try:
            ws = load_workbook(path, read_only=True, data_only=True).worksheets[0]
        except Exception as exc:
            raise BatchFileError(f"Не удалось прочитать Excel: {exc}") from None
        return [
            ["" if v is None else _cell(v) for v in row] for row in ws.iter_rows(values_only=True)
        ]
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "cp1251"):
        try:
            text = raw.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    else:  # pragma: no cover
        raise BatchFileError("Не удалось определить кодировку CSV")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=";,\t")
    except csv.Error:
        dialect = _Semicolon  # русский Excel сохраняет CSV через «;»
    return [row for row in csv.reader(io.StringIO(text), dialect)]


class _Semicolon(csv.excel):
    delimiter = ";"


def _cell(value) -> str:
    if isinstance(value, float):
        return repr(value).removesuffix(".0") if value.is_integer() else repr(value)
    return str(value)


@dataclass(frozen=True, slots=True)
class BatchRow:
    line: int  # номер строки в файле (с заголовком — 1)
    name: str
    subject: str
    values: dict[str, Decimal]  # заполненные числовые колонки
    errors: tuple[str, ...] = ()


def _dec(text: str) -> Decimal:
    clean = str(text).strip().replace("\xa0", "").replace(" ", "").replace(",", ".")
    clean = clean.removesuffix("%").removesuffix("₽")
    return Decimal(clean)


def parse_rows(table: list[list[str]]) -> list[BatchRow]:
    """Заголовок — первая непустая строка; ошибки в строках — у самих строк."""
    rows = [list(r) for r in table]
    start = next((i for i, r in enumerate(rows) if any(str(c).strip() for c in r)), None)
    if start is None:
        raise BatchFileError("Файл пустой")
    header = rows[start]
    columns: dict[int, str] = {}
    for i, title in enumerate(header):
        key = _ALIASES.get(_norm(title))
        if key and key not in columns.values():
            columns[i] = key
    missing = [f.header for f in FIELDS if f.required and f.key not in columns.values()]
    if missing:
        raise BatchFileError(
            "Нет обязательных колонок: " + ", ".join(missing) + ". Скачайте шаблон — в нём "
            "все колонки подписаны."
        )
    out: list[BatchRow] = []
    for n, row in enumerate(rows[start + 1 :], start + 2):
        cells = {key: str(row[i]).strip() if i < len(row) else "" for i, key in columns.items()}
        if not any(cells.values()):
            continue
        values, errors = {}, []
        for key in DECIMAL_KEYS:
            text = cells.get(key, "")
            if not text:
                continue
            try:
                values[key] = _dec(text)
            except InvalidOperation:
                errors.append(f"«{text}» — не число ({_header(key)})")
        for f in FIELDS:
            if (
                f.required
                and f.key in DECIMAL_KEYS
                and f.key not in values
                and not any(e.endswith(f"({f.header})") for e in errors)
            ):
                errors.append(f"не заполнено: {f.header}")
        name = cells.get("name") or f"Строка {n}"
        out.append(BatchRow(n, name, cells.get("subject", ""), values, tuple(errors)))
    if not out:
        raise BatchFileError("В файле нет строк с товарами")
    return out


def _header(key: str) -> str:
    return next(f.header for f in FIELDS if f.key == key)


@dataclass(frozen=True, slots=True)
class BatchResult:
    row: BatchRow
    result: WbResult | None = None
    other: WbResult | None = None  # та же цена, другая схема — для сравнения
    error: str | None = None
    commission_source: str = ""  # «файл», «предмет WB», «форма»

    @property
    def ok(self) -> bool:
        return self.result is not None


@dataclass(slots=True)
class BatchSummary:
    total: int = 0
    computed: int = 0
    losses: int = 0
    low_margin: int = 0
    errors: int = 0
    profit_sum: Decimal = field(default_factory=Decimal)


SubjectLookup = Callable[[str], tuple[Decimal | None, Decimal | None] | None]


def build_input(
    row: BatchRow, base: WbInput, subject: SubjectLookup | None = None
) -> tuple[WbInput, str]:
    v = row.values
    changes = {k: v[k] for k in DECIMAL_KEYS if k in v and k != "retail_price"}
    source = "файл" if {"commission_fbo_pct", "commission_fbs_pct"} & v.keys() else "форма"
    if row.subject and subject is not None and (found := subject(row.subject)) is not None:
        fbo, fbs = found
        if "commission_fbo_pct" not in v and fbo is not None:
            changes["commission_fbo_pct"] = fbo
        if "commission_fbs_pct" not in v and fbs is not None:
            changes["commission_fbs_pct"] = fbs
        source = "предмет WB" if source == "форма" else source
    if "retail_price" in v:
        changes |= {"target": Target.PRICE, "retail_price": v["retail_price"]}
    elif base.target is Target.PRICE:
        raise WbInputError("нет цены до скидки (или выберите в форме подбор цены под цель)")
    return replace(base, name=row.name, **changes), source


def evaluate_batch(
    rows: list[BatchRow],
    base: WbInput,
    tariffs: Tariffs | None = None,
    subject: SubjectLookup | None = None,
) -> list[BatchResult]:
    out = []
    for row in rows:
        if row.errors:
            out.append(BatchResult(row, error="; ".join(row.errors)))
            continue
        try:
            inp, source = build_input(row, base, subject)
            result = solve(inp, tariffs)
            other_scheme = Scheme.FBS if inp.scheme is Scheme.FBO else Scheme.FBO
            other = solve(
                replace(inp, scheme=other_scheme, target=Target.PRICE,
                        retail_price=result.retail_price),
                tariffs,
            )  # fmt: skip
        except WbInputError as exc:
            out.append(BatchResult(row, error=str(exc)))
            continue
        out.append(BatchResult(row, result, other, None, source))
    return out


def summarize(results: list[BatchResult], low_margin_pct: Decimal) -> BatchSummary:
    s = BatchSummary(total=len(results))
    for r in results:
        if not r.ok:
            s.errors += 1
            continue
        s.computed += 1
        s.profit_sum += r.result.profit
        if r.result.profit < 0:
            s.losses += 1
        elif r.result.margin is not None and r.result.margin * 100 < low_margin_pct:
            s.low_margin += 1
    return s


def level(r: BatchResult, low_margin_pct: Decimal) -> str:
    """«error», «loss», «low» или «ok» — для подсветки строки."""
    if not r.ok:
        return "error"
    if r.result.profit < 0:
        return "loss"
    if r.result.margin is not None and r.result.margin * 100 < low_margin_pct:
        return "low"
    return "ok"


# --- файлы Excel --------------------------------------------------------------------------

EXAMPLE = (
    ("Кружка керамическая 350 мл", "12", "10", "11", "180", "990", "30", "Кружки", "", "",
     "", "", "15"),
    ("Плед флисовый 150×200", "35", "30", "8", "620", "2490", "40", "Пледы", "", "", "85",
     "5", ""),
)  # fmt: skip


def write_template(path: Path) -> None:
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Товары"
    for c, f in enumerate(FIELDS, 1):
        cell = ws.cell(1, c, f.header)
        cell.font = Font(bold=True)
        note = ("Обязательная колонка. " if f.required else "Пусто — значение из формы. ") + f.hint
        cell.comment = Comment(note.strip(), "Калькулятор ВЭД")
        ws.column_dimensions[cell.column_letter].width = max(12, len(f.header) + 2)
    ws.column_dimensions["A"].width = 32
    for r, row in enumerate(EXAMPLE, 2):
        for c, value in enumerate(row, 1):
            if value:
                ws.cell(r, c, _number_or_text(value))
    wb.save(path)


def _number_or_text(value: str):
    try:
        number = Decimal(value)
    except InvalidOperation:
        return value
    return int(number) if number == number.to_integral_value() else float(number)


RESULT_HEADERS = (
    "Строка", "Название", "Цена до скидки", "Цена продажи", "Комиссия, %", "К перечислению",
    "Прибыль", "Маржа, %", "ROI, %", "Прибыль по другой схеме", "Ошибка",
)  # fmt: skip


def result_values(r: BatchResult) -> tuple:
    if not r.ok:
        return (r.row.line, r.row.name, *([None] * 8), r.error)
    res, other = r.result, r.other
    commission = (
        res.inp.commission_fbo_pct if res.inp.scheme is Scheme.FBO else res.inp.commission_fbs_pct
    )
    pct = lambda x: None if x is None else round(x * 100, 2)  # noqa: E731
    return (
        r.row.line, r.row.name, res.retail_price, res.price, commission, res.payout,
        res.profit, pct(res.margin), pct(res.roi), other.profit if other else None, None,
    )  # fmt: skip


def write_results(results: list[BatchResult], scheme: Scheme, low_margin_pct: Decimal,
                  path: Path) -> None:  # fmt: skip
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    fills = {
        "loss": PatternFill("solid", fgColor="F8D7D3"),
        "low": PatternFill("solid", fgColor="FFF1C7"),
        "error": PatternFill("solid", fgColor="EEEEEE"),
    }
    wb = Workbook()
    ws = wb.active
    ws.title = "Результаты"
    other = "FBS" if scheme is Scheme.FBO else "FBO"
    headers = [h.replace("другой схеме", other) for h in RESULT_HEADERS]
    for c, h in enumerate(headers, 1):
        ws.cell(1, c, h).font = Font(bold=True)
    for r, res in enumerate(results, 2):
        fill = fills.get(level(res, low_margin_pct))
        for c, value in enumerate(result_values(res), 1):
            cell = ws.cell(r, c, value)
            if isinstance(value, Decimal):
                cell.number_format = "#,##0.00"
            if fill is not None:
                cell.fill = fill
    s = summarize(results, low_margin_pct)
    r = len(results) + 3
    for label, value in (("Схема", scheme.label), ("Товаров", s.total), ("Рассчитано", s.computed),
                         ("Убыточных", s.losses),
                         (f"Маржа ниже {low_margin_pct.normalize():f}%", s.low_margin),
                         ("С ошибками", s.errors)):  # fmt: skip
        ws.cell(r, 1, label).font = Font(bold=True)
        ws.cell(r, 2, value)
        r += 1
    for col, width in zip("ABCDEFGHIJK", (8, 34, 14, 14, 12, 16, 12, 10, 10, 18, 40), strict=True):
        ws.column_dimensions[col].width = width
    wb.save(path)
